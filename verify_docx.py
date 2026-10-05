# -*- coding: utf-8 -*-
"""交付前完整自检：名称替换、标题编号连续性、图表编号、占位清单、标点一致性。"""
import re
import sys
import zipfile
from collections import Counter

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")
TARGET = "青阅循环_优化版.docx"
doc = Document(TARGET)
body = doc.element.body

print("=" * 74)
print("1) 全局替换检查")
print("=" * 74)
with zipfile.ZipFile(TARGET) as z:
    total_old = 0
    for n in z.namelist():
        if n.endswith(".xml") or n.endswith(".rels"):
            c = z.read(n).decode("utf-8", "ignore").count("书环智循环")
            if c:
                print(f"  [残留] {n}: {c}")
                total_old += c
    doc_txt = "".join(p.text for p in doc.paragraphs)
print(f"  「书环智循环」全文残留数：{total_old}  （应为 0）")
print(f"  「青阅循环」出现次数（正文段落）：{doc_txt.count('青阅循环')}")

print()
print("=" * 74)
print("2) 标题编号连续性")
print("=" * 74)
heads = []
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Heading 2", "Heading 3"):
        heads.append((p.style.name, p.text.strip()))
chapters = [t for s, t in heads if s == "Heading 1"]
print("  一级标题：", " | ".join(chapters))
nums = Counter()
problems = []
for s, t in heads:
    m = re.match(r"^(\d+(?:\.\d+)*)\s", t)
    if m:
        nums[m.group(1)] += 1
dups = [k for k, v in nums.items() if v > 1]
print(f"  编号总数 {len(nums)}，重复编号：{dups if dups else '无'}")
# 章内小节序号连续性
for ch in ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10"]:
    subs = sorted(int(k.split(".")[1]) for k in nums if re.fullmatch(rf"{ch}\.\d+", k))
    if subs:
        expect = list(range(1, len(subs) + 1))
        flag = "" if subs == expect else "  <== 不连续"
        print(f"  第{ch}章小节：{subs}{flag}")
        if flag:
            problems.append(ch)

print()
print("=" * 74)
print("3) 图表编号（按章编号，从图 1-1 开始）")
print("=" * 74)
figs = []
for ch in body.iterchildren():
    if ch.tag == qn("w:tbl"):
        tb = Table(ch, doc)
        if len(tb.rows) == 1 and len(tb.columns) == 1:
            txt = tb.rows[0].cells[0].text
            m = re.search(r"图 (\d+)-(\d+) (.*?)（待补充数据）", txt)
            if m and "【图表占位" in txt:
                figs.append((int(m.group(1)), int(m.group(2)), m.group(3)))
print(f"  图注总数：{len(figs)}")
ok = True
by_ch = {}
for c, n, title in figs:
    by_ch.setdefault(c, []).append(n)
for c in sorted(by_ch):
    expect = list(range(1, len(by_ch[c]) + 1))
    flag = "" if by_ch[c] == expect else "  <== 编号不连续"
    if flag:
        ok = False
    print(f"  第{c:>2}章：图 {c}-{by_ch[c]} (共{len(by_ch[c])}个){flag}")
print(f"  编号连续性：{'通过' if ok and figs else '存在问题'}")

print()
print("=" * 74)
print("4) 图表占位清单")
print("=" * 74)
print(f"  {'图号':<8}{'建议图表类型':<12}数据来源 / 标题")
n_ph = 0
for ch in body.iterchildren():
    if ch.tag == qn("w:tbl"):
        t = Table(ch, doc)
        if len(t.rows) == 1 and len(t.columns) == 1:
            txt = t.rows[0].cells[0].text
            if txt.strip().startswith("【图表占位"):
                n_ph += 1
                m = re.search(r"图 ([\d-]+)：(.*?)｜建议图表类型：(.*?)｜数据来源：(.*?)】", txt)
                if m:
                    print(f"  图 {m.group(1):<6}{m.group(3):<12}{m.group(4)}"
                          f"  —— {m.group(2)}")
print(f"  占位表格总数：{n_ph}")

print()
print("=" * 74)
print("5) 中英文标点一致性")
print("=" * 74)
alltxt = "\n".join(p.text for p in doc.paragraphs)
straight_d = alltxt.count('"')
full_d = alltxt.count("\u201c") + alltxt.count("\u201d")
print(f"  直引号 \" 数量：{straight_d}    中文弯引号 “” 数量：{full_d}")
mixed = [p.text[:60] for p in doc.paragraphs
         if '"' in p.text and ("\u201c" in p.text or "\u201d" in p.text)]
print(f"  同一段内混用中英引号的段落数：{len(mixed)}")
for m in mixed[:5]:
    print("    ", m)
cjk_ascii_comma = len(re.findall(r"[\u4e00-\u9fff],", alltxt))
cjk_ascii_period = len(re.findall(r"[\u4e00-\u9fff]\.", alltxt))
print(f"  中文后紧跟半角逗号/句点：{cjk_ascii_comma} / {cjk_ascii_period}")

print()
print("=" * 74)
print("6) 结构概览")
print("=" * 74)
print(f"  节数：{len(doc.sections)}   段落：{len(doc.paragraphs)}   表格：{len(doc.tables)}")
for i, s in enumerate(doc.sections):
    print(f"  节{i}: 页边距 上{s.top_margin.cm:.2f} 下{s.bottom_margin.cm:.2f} "
          f"左{s.left_margin.cm:.2f} 右{s.right_margin.cm:.2f} "
          f"页眉{s.header_distance.cm:.2f} 页脚{s.footer_distance.cm:.2f} "
          f"页眉表{len(s.header.tables)} 页脚段{[p.text for p in s.footer.paragraphs]}")
print(f"  旧名称残留（XML）: {total_old}")
print(f"  编号问题章节: {problems if problems else '无'}")
