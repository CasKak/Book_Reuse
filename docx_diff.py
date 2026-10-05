# -*- coding: utf-8 -*-
"""文档结构差异对比工具。

用途：对比两个 DOCX 的块级结构（段落 / 表格 / 节），列出差异，便于确认
      优化脚本的输出是否与预期一致，或排查他人手工改动。

用法：
    python compare_docx.py <基准文件.docx> <对比文件.docx>
    python compare_docx.py 青阅循环_优化版_用户版备份.docx 青阅循环_优化版.docx
"""
from __future__ import annotations

import difflib
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def outline(path):
    """把文档拍平成 [(类型, 样式/尺寸, 文本)] 列表。"""
    doc = Document(path)
    items = []
    for ch in doc.element.body.iterchildren():
        tag = ch.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(ch, doc)
            txt = p.text.strip()
            if txt:
                items.append(("P", p.style.name, txt))
        elif tag == "tbl":
            t = Table(ch, doc)
            items.append(("T", f"{len(t.rows)}x{len(t.columns)}",
                          t.rows[0].cells[0].text.strip().split("\n")[0][:70]))
        elif tag == "sdt":
            items.append(("SDT", "目录", "目录内容控件"))
        elif tag == "sectPr":
            items.append(("SECTPR", "", "文档节属性"))
        else:
            items.append((tag.upper(), "", ""))
    return doc, items


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    a_path, b_path = sys.argv[1], sys.argv[2]

    da, a = outline(a_path)
    db, b = outline(b_path)

    print(f"基准：{a_path}")
    print(f"  段落 {len(da.paragraphs)}  表格 {len(da.tables)}  节 {len(da.sections)}")
    print(f"对比：{b_path}")
    print(f"  段落 {len(db.paragraphs)}  表格 {len(db.tables)}  节 {len(db.sections)}")

    def fmt(it):
        return f"{it[0]}|{it[1]}|{it[2]}"

    diff = list(difflib.unified_diff(
        [fmt(x) for x in a], [fmt(x) for x in b],
        fromfile=a_path, tofile=b_path, lineterm="", n=1))

    print()
    if not diff:
        print("结构完全一致。")
    else:
        changed = [d for d in diff if d[:1] in "+-"]
        print(f"结构差异（{len(changed)} 行）：")
        for line in diff:
            print("  " + line)


if __name__ == "__main__":
    main()
